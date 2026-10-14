#!/usr/bin/env python3
"""Build matching editable PPTX and vector PDF slides, plus review images.

Dependencies: reportlab, python-pptx, Pillow; pdftoppm for image rendering.
All diagram shapes and text remain editable in the PowerPoint file.
"""
from __future__ import annotations
import hashlib
import json
import subprocess
from pathlib import Path

from PIL import Image, ImageOps, ImageDraw
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
W,H = 1280,720
INK='#152A3B'; MUTED='#506371'; TEAL='#087F83'; MINT='#E9F5F1'
BLUE='#ECF2F8'; LIGHT='#F5F7F9'; LINE='#D6E0E5'; AMBER='#AA5C20'
AMBER_LIGHT='#FFF2E5'; WHITE='#FFFFFF'
for key,file in [('Deck','DejaVuSans.ttf'),('DeckBold','DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(key, '/usr/share/fonts/truetype/dejavu/'+file))

class Deck:
    def __init__(self):
        self.pdf=canvas.Canvas(str(OUT/'FinNLP65ADMLSlides.pdf'),pagesize=(960,540),
            invariant=1,initialFontName='Deck',initialFontSize=24)
        self.pdf.setTitle(json.loads((HERE/'script.json').read_text())['title'])
        self.pdf.setAuthor('Gautier Marti; Khaled AlNuaimi')
        self.ppt=Presentation()
        self.ppt.slide_width=Inches(W/96);self.ppt.slide_height=Inches(H/96)
        self.ppt.core_properties.title=json.loads((HERE/'script.json').read_text())['title']
        self.ppt.core_properties.author='Gautier Marti; Khaled AlNuaimi'
        self.bounds=[];self.n=0
    def new(self,dark=False):
        if self.n:self.pdf.showPage()
        self.n+=1
        self.pdf.scale(.75,.75)
        self.slide=self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        self.rect(0,0,W,H,INK if dark else WHITE)
        c='#B6CFD6' if dark else MUTED
        self.text(54,684,'ADML   ·   FinNLP 2026   ·   FinNLP-65',15,color=c)
        self.text(1158,684,f'{self.n:02d} / 06',15,color=c)
    def rect(self,x,y,w,h,color,r=0):
        self.pdf.setFillColor(HexColor(color));self.pdf.setStrokeColor(HexColor(color))
        self.pdf.roundRect(x,H-y-h,w,h,r,stroke=0,fill=1)
        sh=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,
            Inches(x/96),Inches(y/96),Inches(w/96),Inches(h/96))
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color[1:]);sh.line.fill.background()
        if r:sh.adjustments[0]=min(.1,r/min(w,h))
    def text(self,x,y,t,size=28,bold=False,color=INK):
        font='DeckBold' if bold else 'Deck'
        width=pdfmetrics.stringWidth(t,font,size)
        assert x>=30 and x+width<=1250,(self.n,t,x+width)
        assert y>=0 and y+size*1.3<=715,(self.n,t,y)
        self.bounds.append(dict(slide=self.n,text=t,x=x,y=y,width=width,size=size))
        self.pdf.setFont(font,size);self.pdf.setFillColor(HexColor(color))
        self.pdf.drawString(x,H-y-pdfmetrics.getAscent(font,size),t)
        tx=self.slide.shapes.add_textbox(Inches(x/96),Inches(y/96),Inches((width+16)/96),Inches(size*1.45/96))
        tf=tx.text_frame;tf.margin_left=tf.margin_top=tf.margin_right=tf.margin_bottom=0
        tf.word_wrap=False
        p=tf.paragraphs[0];p.text=t;p.font.name='DejaVu Sans';p.font.size=Pt(size*.75)
        p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color[1:])
    def para(self,x,y,t,width,size=26,bold=False,color=INK,leading=None):
        lines=[];font='DeckBold' if bold else 'Deck'
        for para in t.split('\n'):
            line=''
            for word in para.split():
                candidate=(line+' '+word).strip()
                if pdfmetrics.stringWidth(candidate,font,size)>width and line:
                    lines.append(line);line=word
                else:line=candidate
            lines.append(line)
        lead=leading or size*1.33
        for i,line in enumerate(lines):self.text(x,y+i*lead,line,size,bold,color)
        return y+len(lines)*lead
    def header(self,kicker,title):
        self.text(54,37,kicker.upper(),18,True,TEAL)
        self.text(54,81,title,40,True)
        self.rect(54,143,1172,3,LINE)
    def notes(self,n):
        s=json.loads((HERE/'script.json').read_text())['slides'][n-1]
        self.slide.notes_slide.notes_text_frame.text='\n\n'.join(s['paragraphs'])
    def save(self):
        self.pdf.save();self.ppt.save(OUT/'FinNLP65ADMLSlides.pptx')
        (HERE/'slide_text.json').write_text(json.dumps(self.bounds,indent=2)+'\n')

def build():
    d=Deck()
    d.new(True)
    d.text(54,36,'ACCEPTED PAPER · FinNLP-65 · POSTER PRESENTATION',18,True,'#70D9CD')
    d.text(54,91,'ADML at the FinEval 2026:',40,True,WHITE)
    d.text(54,146,'Version-Aware Retrieval and Grounded Sparse',36,True,WHITE)
    d.text(54,194,'Editing for Multilingual Financial QA',36,True,WHITE)
    d.text(54,268,'Agentic research across three financial NLP tasks',28,False,'#A4E5DB')
    d.text(54,329,'Gautier Marti¹  ·  Khaled AlNuaimi¹ ²',25,False,WHITE)
    d.text(54,370,'¹ ADIA     ² Khalifa University',20,False,'#B6CFD6')
    for x,rank,label,score in [(54,'1st','Greek financial NER','Entity F1  0.9676'),(454,'Tied 1st','Japanese ICR','Accuracy  0.9714'),(854,'2nd','Multilingual QA','ROUGE-1 F1  0.3726')]:
        d.rect(x,433,372,163,'#213D50',10)
        d.text(x+24,451,rank,42,True,'#70D9CD')
        d.text(x+24,511,label,24,True,WHITE)
        d.text(x+24,550,score,21,False,'#D6E8EB')
    d.text(54,622,'Official private results · Different task metrics',18,False,'#B6CFD6')
    d.text(968,622,'Synthetic narration',18,False,'#B6CFD6');d.notes(1)

    d.new();d.header('Agentic research method','How we ran the experiments')
    stages=[('01','Audit','Data + metric'),('02','Hypothesize','Target a failure'),('03','Implement','Controlled trial'),('04','Evaluate','Transfer + harm'),('05','Decide','Promote / reject')]
    for i,(n,a,b) in enumerate(stages):
        x=54+i*239
        d.rect(x,183,216,153,MINT if i==4 else BLUE,9)
        d.text(x+18,202,n,20,True,TEAL)
        d.text(x+18,242,a,28,True)
        d.text(x+18,287,b,18,color=MUTED)
        if i<4:d.text(x+221,244,'›',27,True,TEAL)
    d.text(54,365,'Evaluate across entity groups, sources, companies, and time.',27,True)
    d.text(54,412,'Record corrections, new errors, negative results, and revised decisions.',24,color=MUTED)
    for x,n,label in [(54,'209','research scripts'),(352,'32','source modules'),(650,'99','test files'),(948,'25','preregistrations')]:
        d.rect(x,477,278,120,LIGHT,8)
        d.text(x+20,490,n,43,True,TEAL);d.text(x+20,550,label,22)
    d.text(54,622,'Repository snapshot 557bf38 · 48,468 physical Python lines, including tests',17,color=MUTED)
    d.text(54,647,'Research code and experiment records',17,color=MUTED);d.notes(2)

    d.new();d.header('Systems designed for new inputs','Models we tried and changes we kept')
    cols=[(54,'Greek NER','GreekBERT · legal encoders\nXLM-R · GLiNER','Numeric processing\nSpan consensus\nConservative repairs',MINT),
          (454,'Japanese ICR','BERT · Rinna RoBERTa\nModernBERT · Nemotron 9B','Domain adaptation\nOrdinal decoding\nAdditional supervision',BLUE),
          (854,'Multilingual QA','Eva 4B · Qwen 27B / 35B\nQwen 35B / 57B planners','Typed financial facts\nArithmetic + provenance\nGrounded sparse edits',MINT)]
    for x,title,models,methods,col in cols:
        d.rect(x,176,372,281,col,10)
        d.text(x+22,197,title,29,True,TEAL)
        d.para(x+22,248,models,328,21,leading=31)
        d.para(x+22,337,methods,328,23,leading=34)
    d.rect(54,477,1172,103,AMBER_LIGHT,9)
    d.text(75,493,'REJECTED',17,True,AMBER)
    d.text(75,527,'ModernBERT: 5/5 source folds worse',24,True)
    d.text(665,527,'Qwen filter: one issuer exceeds error limit',22,True)
    d.text(54,604,'A positive result: Greek consensus +4 corrections / 0 new errors (80 rows).',23,True,TEAL)
    d.text(54,641,'Greek validation repeatedly inspected; each task used its own validation protocol.',17,color=MUTED);d.notes(3)

    d.new();d.header('QA case study · 192 historical examples','Choosing which answers to edit')
    d.text(54,177,'ROUGE-1 F1',19,True,MUTED)
    bars=[('Grounded baseline',.3652420,TEAL),('Reference-aware oracle',.4029502,AMBER),('Nested held-out selector',.3652420,TEAL)]
    for i,(label,value,color) in enumerate(bars):
        y=218+i*105
        d.text(54,y,label,23,True)
        d.rect(54,y+38,582,29,LIGHT,4)
        d.rect(54,y+38,582*value/.45,29,color,4)
        d.text(655,y+30,f'{value:.4f}',26,True,color)
    d.rect(818,190,408,334,BLUE,10)
    d.para(843,215,'Useful edits existed.',346,29,True)
    d.para(843,268,'Choosing edits for a new company remained difficult.',346,27)
    d.para(843,397,'Nested selector: no edits.\nTemporal test: 4 edits,\nall 4 harmful.',346,23,color=MUTED,leading=33)
    d.rect(54,563,1172,66,AMBER_LIGHT,8)
    d.text(76,582,'Oracle: uses reference answers unavailable at inference.',26,True,AMBER)
    d.text(54,645,'Historical evaluation · Not directly comparable with the official private leaderboard.',17,color=MUTED);d.notes(4)

    d.new();d.header('Final competition systems','What the benchmark ultimately rewarded')
    for x,title,body,bottom,col in [(54,'Greek NER','200 / 200 exact inputs','Public-reference retrieval',MINT),
        (454,'Japanese ICR','Explicit intent cues','Deterministic semantic rubric',BLUE),
        (854,'Multilingual QA','76 / 76 exact inputs','Public-reference retrieval',MINT)]:
        d.rect(x,177,372,159,col,9)
        d.text(x+20,194,title,28,True,TEAL)
        d.text(x+20,245,body,23,True)
        d.text(x+20,289,bottom,20,color=MUTED)
    d.text(54,361,'External data permitted · Final runner: CPU only, no neural model',24,True)
    d.rect(54,417,520,181,LIGHT,10)
    d.text(77,435,'QA PUBLIC',18,True,MUTED);d.text(335,435,'QA PRIVATE',18,True,MUTED)
    d.text(77,475,'0.9994',46,True,TEAL);d.text(335,475,'0.3726',46,True,AMBER)
    d.text(77,544,'24 examples',21,color=MUTED);d.text(335,544,'52 examples',21,color=MUTED)
    d.para(627,427,'Same inputs.\nDifferent target versions?',551,33,True,leading=47)
    d.para(627,539,'Consistent with reference changes;\nprivate reference answers unavailable.',551,23,color=MUTED,leading=33)
    d.text(54,633,'QA scores: official ROUGE-1 F1 · Retrieval coverage refers to full-query input matching.',17,color=MUTED);d.notes(5)

    d.new(True)
    d.text(54,37,'LESSONS FROM THREE TASKS',18,True,'#70D9CD')
    d.para(54,89,'What we learned',1164,49,True,WHITE,leading=65)
    for y,n,a,b in [(269,'01','Check data and references','Look for repeated inputs and changed answers.'),
        (375,'02','Test on new sources','Track corrections, new errors, and unchanged answers.'),
        (481,'03','Keep failed experiments','Use validation results to decide what to keep.')]:
        d.text(55,y,n,32,True,'#70D9CD')
        d.text(132,y,a,31,True,WHITE);d.text(132,y+46,b,25,color='#C7DDDF')
    d.text(54,605,'The final benchmark left the broader model work unresolved.',21,color='#B6CFD6')
    d.text(54,645,'Gautier Marti · Khaled AlNuaimi     |     Thank you',22,True,WHITE);d.notes(6)
    d.save()

    slides=HERE/'slides';slides.mkdir(exist_ok=True)
    subprocess.run(['pdftoppm','-r','144','-png',str(OUT/'FinNLP65ADMLSlides.pdf'),str(slides/'slide')],check=True)
    imgs=sorted(slides.glob('slide-*.png'))
    sheet=Image.new('RGB',(1320,1184),'#E3E9EE')
    for i,path in enumerate(imgs):
        im=Image.open(path).convert('RGB').resize((640,360),Image.Resampling.LANCZOS)
        sheet.paste(im,(10+(i%2)*660,10+(i//2)*394))
        ImageDraw.Draw(sheet).text((18+(i%2)*660,374+(i//2)*394),f'Slide {i+1}',fill=INK)
    sheet.save(OUT/'FinNLP65ADMLSlidesPreview.jpg',quality=93)
    script=json.loads((HERE/'script.json').read_text())
    text=['# FinNLP-65 narration script','',script['title'],'','Synthetic narration. '+str(sum(len(p.split()) for s in script['slides'] for p in s['paragraphs']))+' words.','']
    for s in script['slides']:
        text+=['## Slide '+str(s['number'])+' — '+s['title'],'']
        for p in s['paragraphs']:text+=[p,'']
    (HERE/'script.md').write_text('\n'.join(text))
    print('Built PDF, editable PPTX, 6 slide images, preview and script.')

if __name__=='__main__':build()
